# -*- coding: utf-8 -*-
"""生成《智愈寻真》项目计划书 Word 文档（报名用）"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# ---- 全局中文字体 ----
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_cn(run, font='宋体'):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)

def heading(text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    set_cn(run, '黑体')
    return h

def para(text, bold=False, size=11, align=None, font='宋体'):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_cn(run, font)
    if align:
        p.alignment = align
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    set_cn(run, '宋体')
    return p

# ============ 封面 ============
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('项目计划书')
r.font.size = Pt(28); r.bold = True
set_cn(r, '黑体')

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('智愈寻真 —— 内科教研协同智能体平台')
r.font.size = Pt(16)
set_cn(r, '黑体')

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run('参赛赛道：学科垂类大模型与创新应用开发大赛（助教 / 助学）')
r.font.size = Pt(12)
set_cn(r, '宋体')

doc.add_paragraph()

# ============ 一、项目概述 ============
heading('一、项目概述', 1)
para('智愈寻真是面向高校一流内科学建设的"教—学—管"三位一体人工智能训练平台，'
     '基于学科垂类大模型、检索增强生成与多智能体编排技术，为医学院校学生与临床教师'
     '提供覆盖教学、练习、批阅、学情分析全流程的智能化支持。平台以人工智能中台为核心，'
     '融合深度推理、可信知识库检索与智能体协作，构建人工智能知识问答与讲解、智能备课、'
     '学情预警、简答及论述主观题批改四大核心能力。')
para('项目愿景是构建"人机协同"的内科医学教育新范式，让优质教学资源突破时间与空间限制，'
     '在低风险环境中高频训练临床思维，同时把教师从重复性批阅工作中释放出来，聚焦个性化教学。')

# ============ 二、项目背景与行业痛点 ============
heading('二、项目背景与行业痛点', 1)
para('高校临床医学教育正从"知识讲授"转向"能力训练"。内科学是临床医学主干课程，'
     '覆盖疾病谱广、诊断链条长、鉴别诊断复杂，是学生从基础医学走向临床实战时最易出现断层的地方。'
     '当前教学现场存在三个突出矛盾：')
bullet('学生练习机会不足：真人标准化病人培训与组织成本高，学生难以高频、低压力地练习问诊、鉴别诊断和临床沟通。')
bullet('教师批阅负担过重：大病历批改耗时长，教师通常需逐份阅读、标注逻辑漏洞、纠正错漏，难以兼顾临床与教学。')
bullet('教学管理缺少数据闭环：教研室知道学生"考得不好"，却难以定位到"哪类疾病、哪一步问诊、哪个检查决策"持续薄弱。')
para('大模型、多模态识别、知识库检索与智能体编排技术，已具备将虚拟病人训练、临床思维评估、'
     '作业批阅与学情分析打通的条件。本项目目标就是把人工智能从简单问答工具，升级为面向内科学建设的教研协同平台。')

# ============ 三、产品定位与目标用户 ============
heading('三、产品定位与目标用户', 1)
para('产品定位：面向高校一流内科学建设的"教—学—管"三位一体人工智能训练平台。')
para('目标用户与关键诉求：')
t = doc.add_table(rows=1, cols=3)
t.style = 'Light Grid Accent 1'
hdr = t.rows[0].cells
for i, txt in enumerate(['用户类型', '覆盖范围', '关键诉求']):
    hdr[i].text = txt
    for p in hdr[i].paragraphs:
        for run in p.runs:
            set_cn(run, '黑体'); run.bold = True
rows = [
    ('主体学生', '大三、大四临床医学生', '将教材知识转化为问诊逻辑、诊断路径与病例书写能力'),
    ('辐射学生', '大二至大五及规培生', '基础衔接、每日训练、接近真实的客观结构化临床考试训练'),
    ('教师用户', '临床带教主治医师、教学秘书', '快速出题、批阅复核、查看班级薄弱点、复用优质病例'),
    ('管理用户', '教研室主任、管理员、运维', '看全局数据、管用户、管内容、管模型成本与安全合规'),
]
for u, c, n in rows:
    cells = t.add_row().cells
    cells[0].text = u; cells[1].text = c; cells[2].text = n
    for cell in cells:
        for p in cell.paragraphs:
            for run in p.runs:
                set_cn(run, '宋体')

# ============ 四、核心功能 ============
heading('四、核心功能（四大核心能力）', 1)

heading('4.1 人工智能知识问答与讲解', 2)
para('学生端提供沉浸式问诊训练与教材问答：可上传影像、检验图片并圈画异常点，人工智能病人以'
     '流式对话方式回复，模拟真实交流节奏；右侧临床思维决策树实时展示已问内容、已排除诊断、'
     '遗漏检查与检查成本。知识问答基于可信教材知识库检索，回答均附书名、章节、页码溯源，'
     '降低幻觉风险。')

heading('4.2 智能备课', 2)
para('教师端提供零代码标准化病人配置台，通过表单快速创建虚拟病例（主诉、隐藏诊断、性格风格、'
     '检查项目与费用、标准路径），并支持课件（文档、演示、音视频）独立下发。'
     '病例可发布至病例广场，供同行预览、引用与评分；引用时生成独立副本，原病例更新不影响已下发作业。'
     '支持将标准化病人与病例大厅引用到作业，快速完成备课。')

heading('4.3 学情预警', 2)
para('系统将班级训练行为与人工智能评估结果转化为可行动的教学洞察：作业完成率、平均客观结构化'
     '临床考试分、共性漏问项、共性误诊、过度检查、薄弱知识点。学情数据通过站内信推送预警，'
     '不引入外部消息通道，帮助教师及时定位并干预班级共性薄弱点，形成"训练—预警—补救"闭环。')

heading('4.4 简答 / 论述主观题批改', 2)
para('超越传统大病历批阅，支持教师自定义评分要点，对简答、论述等主观题做结构化批改：从文书规范、'
     '医学事实、逻辑链条、鉴别诊断、人文表达多维度给出错误高亮、扣分项与综合评语。'
     '批改前先执行"格式盾牌"结构与必填项校验，不通过则打回学生修改；教师可复核并覆盖人工智能结果，'
     '最终成绩以教师为准，所有覆盖操作留痕可审计。')

# ============ 五、技术方案与系统架构 ============
heading('五、技术方案与系统架构', 1)
para('系统采用异构服务架构，职责清晰、可独立演进：')
bullet('移动端（Flutter）：学生与教师共用同一应用包，登录后按角色分流到各自首页；负责问诊、配置、批阅等核心教学功能。')
bullet('管理端（Vue 3 Web）：全局驾驶舱、用户权限、内容风控、系统配置与审计日志。')
bullet('业务中台（Spring Boot）：鉴权、用户、班级、病例、作业、批阅、审计与文件网关，是业务数据的唯一事实来源。')
bullet('人工智能中台（FastAPI）：大模型调度、多智能体编排、知识库检索、流式输出。')
bullet('数据底座：MySQL 存储业务数据，Milvus 存储教材切片向量，对象存储存放教材与报告。')

para('人工智能中台内置多智能体分工：标准化病人智能体扮演虚拟病人；导师智能体维护思维树并触发苏格拉底式引导；'
     '视觉智能体处理影像与圈选区域；评估智能体生成客观结构化临床考试四维评分与复盘；'
     '批阅智能体批改主观题；安全智能体识别越界请求并截断。关键链路遵循"规则优先、绝不伪造结果"原则：'
     '能由规则确定的不交给模型，人工智能不可用时诚实降级并明确告知，评分判读等关键结果绝不输出假分数。')

# ============ 六、核心创新点 ============
heading('六、核心创新点', 1)
bullet('学科垂类大模型 + 可信知识库：基于授权教材构建向量知识库，所有人工智能回答附章节页码溯源，医学内容经专家审核。')
bullet('临床思维可视化：将抽象的诊断推理过程实时呈现为可交互的思维树，让学生看见漏问、误排除与过度检查。')
bullet('多智能体协作的人机协同：以标准化病人为唯一自主智能体，其余能力收敛为可重放、可复核的工作流，兼顾灵活与可控。')
bullet('教学闭环数据化：从训练、批改到学情预警形成完整数据闭环，把教师经验沉淀为可复用的班级补救建议。')
bullet('安全可信的工程伦理：规则优先、诚实降级、全链路追踪与审计，使人工智能训练平台可靠、可信、好用。')

# ============ 七、实施计划与里程碑 ============
heading('七、实施计划与里程碑', 1)
t2 = doc.add_table(rows=1, cols=3)
t2.style = 'Light Grid Accent 1'
hdr2 = t2.rows[0].cells
for i, txt in enumerate(['阶段', '目标', '交付物']):
    hdr2[i].text = txt
    for p in hdr2[i].paragraphs:
        for run in p.runs:
            set_cn(run, '黑体'); run.bold = True
phases = [
    ('阶段一：底座与核心闭环', '建立生产级底座，跑通问诊—批阅—学情主线', '统一登录、权限与审计；智能体问诊与批改；学情看板'),
    ('阶段二：完整教学产品', '补齐学生增长与教师复用能力', '病例广场、每日一例、错题本、人工智能复盘报告'),
    ('阶段三：治理与规模化', '支撑学院级真实使用与长期运营', '驾驶舱、内容风控、模型容灾、成本监控'),
    ('阶段四：试点上线', '真实教学场景灰度上线并持续优化', '试点班级、反馈周报、模型与病例持续迭代'),
]
for s, g, d in phases:
    cells = t2.add_row().cells
    cells[0].text = s; cells[1].text = g; cells[2].text = d
    for cell in cells:
        for p in cell.paragraphs:
            for run in p.runs:
                set_cn(run, '宋体')

para('关键效能指标（示例目标）：单份大病历批阅时长缩短约七成、多模态对话首词延迟不超过一点五秒、'
     '试点班级客观结构化临床考试"问诊逻辑"单项均分提升不少于百分之十五。')

# ============ 八、团队与资源保障 ============
heading('八、团队与资源保障', 1)
para('项目由医学教育背景的产品与研发团队协作推进，配备临床教学顾问把关医学准确性，'
     '技术栈覆盖 Flutter 移动端、Vue 管理端、Spring Boot 业务中台与 FastAPI 人工智能中台。'
     '开发环境已搭建国内镜像加速的构建链路，后端依赖经容器化部署，移动端支持真机与模拟器双通道调试。')

# ============ 九、预期成果与价值 ============
heading('九、预期成果与价值', 1)
bullet('对学生：高频、低压力的虚拟病人训练与即时复盘，缩短临床思维培养周期。')
bullet('对教师：备课与批阅效率显著提升，从重复劳动中释放，聚焦个性化指导。')
bullet('对院校：沉淀可复用的病例与教学资源，形成数据驱动的教学治理闭环。')
bullet('对赛道：提供"学科垂类大模型 + 助教助学"可落地范式，具备推广与持续运营价值。')

# ============ 十、风险与应对 ============
heading('十、风险与应对', 1)
bullet('人工智能医学回答偏差：强制知识库溯源、医学专家审核病例、人工智能输出免责声明、教师最终复核。')
bullet('模型成本不可控：设置用量预算与阈值预警、批阅队列限流、缓存常用教材召回。')
bullet('教师信任不足：保留人工复核与覆盖机制，展示扣分依据，支持申诉。')
bullet('隐私与合规：权限隔离、敏感信息脱敏、审计日志、导出审批，训练数据不接真实患者信息。')

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run('—— 智愈寻真项目组 ——')
r.font.size = Pt(10)
set_cn(r, '宋体')

out = r'E:\zhiyu\docs\项目计划书_智愈寻真.docx'
doc.save(out)
print('Saved:', out)
